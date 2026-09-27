import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palet Warna
    NAVY = RGBColor(15, 23, 42)        # #0F172A
    SLATE_800 = RGBColor(30, 41, 59)   # #1E293B
    SLATE_100 = RGBColor(241, 245, 249)# #F1F5F9
    WHITE = RGBColor(255, 255, 255)
    BLUE = RGBColor(37, 99, 235)       # #2563EB
    BLUE_LIGHT = RGBColor(239, 246, 255)
    INDIGO = RGBColor(79, 70, 229)     # #4F46E5
    EMERALD = RGBColor(16, 185, 129)   # #10B981
    EMERALD_LIGHT = RGBColor(236, 253, 245)
    AMBER = RGBColor(245, 158, 11)     # #F59E0B
    AMBER_LIGHT = RGBColor(254, 243, 199)
    RED = RGBColor(239, 68, 68)        # #EF4444
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    BORDER_COLOR = RGBColor(226, 232, 240)

    def add_header(slide, title_text, category_text="PROMAP V2.4 — WORKFLOW & ROLE RBAC"):
        # Header category badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.7))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

        # Subtle bottom line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.color.rgb = BORDER_COLOR

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # =========================================================================
    # SLIDE 1: COVER (Dark Navy)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Accent decorative pill
    pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(2.8), Inches(0.4))
    pill.fill.solid()
    pill.fill.fore_color.rgb = SLATE_800
    pill.line.color.rgb = BLUE
    p_pill = pill.text_frame.paragraphs[0]
    p_pill.text = "ENTERPRISE WORKSPACE V2.4"
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = RGBColor(96, 165, 250)
    p_pill.alignment = PP_ALIGN.CENTER

    # Main Title
    t1_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.0), Inches(2.0))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Panduan Alur Kerja & Peran Pengguna"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    
    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Sistem Manajemen Eksekusi Action Plan & Akuntabilitas Multi-Divisi"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = RGBColor(148, 163, 184)
    p1_sub.space_before = Pt(14)

    # 4 Key Role Badges on Cover
    roles = [
        ("SUPER ADMIN", "Multi-Tenant & Kuota"),
        ("ADMIN OPS", "Perusahaan & Divisi"),
        ("MANAGER", "Perencanaan & Approval"),
        ("PIC / EKSEKUTOR", "Eksekusi & Evidence")
    ]
    card_w = Inches(2.65)
    for idx, (r_name, r_desc) in enumerate(roles):
        c_left = Inches(1.0) + idx * (card_w + Inches(0.24))
        card = add_card(slide1, c_left, Inches(4.8), card_w, Inches(1.6), SLATE_800, RGBColor(51, 65, 85))
        tf_c = card.text_frame
        tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_c1 = tf_c.paragraphs[0]
        p_c1.text = r_name
        p_c1.font.size = Pt(13)
        p_c1.font.bold = True
        p_c1.font.color.rgb = RGBColor(56, 189, 248)
        p_c1.alignment = PP_ALIGN.CENTER
        
        p_c2 = tf_c.add_paragraph()
        p_c2.text = r_desc
        p_c2.font.size = Pt(10)
        p_c2.font.color.rgb = RGBColor(203, 213, 225)
        p_c2.alignment = PP_ALIGN.CENTER
        p_c2.space_before = Pt(6)

    # Footer note
    ft1 = slide1.shapes.add_textbox(Inches(1.0), Inches(6.8), Inches(11.0), Inches(0.4))
    p_ft1 = ft1.text_frame.paragraphs[0]
    p_ft1.text = "PT Sobat UMKM Pro • ProMaP Enterprise Execution Platform"
    p_ft1.font.size = Pt(10)
    p_ft1.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # SLIDE 2: LATAR BELAKANG & FILOSOFI PROMAP
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Mengapa ProMaP? Solusi Eksekusi Rencana Strategis", "FILOSOFI & MASALAH BISNIS")

    # Kolom Kiri: Masalah Klasik Perusahaan
    c_prob = add_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), RGBColor(254, 242, 242), RGBColor(254, 202, 202))
    tf_pr = c_prob.text_frame
    tf_pr.margin_left = tf_pr.margin_top = tf_pr.margin_right = tf_pr.margin_bottom = Inches(0.4)
    p_pr1 = tf_pr.paragraphs[0]
    p_pr1.text = "Tantangan Klasik Eksekusi Rencana"
    p_pr1.font.size = Pt(16)
    p_pr1.font.bold = True
    p_pr1.font.color.rgb = RED

    items_prob = [
        "Rencana Strategis Berhenti di Notula: Program kerja hanya hidup saat rapat, tanpa tindak lanjut operasional.",
        "Spreadsheet Berceceran & Tidak Sinkron: Progres tercecer di berbagai file Excel dan grup chat tanpa kendali versi.",
        "Ketiadaan Pembuktian Nyata (No Evidence): Status sering diklaim 'Selesai' tanpa ada verifikasi hasil kerja riil.",
        "Tenggat Waktu Sering Terlewat (Overdue Ghosting): Keterlambatan baru disadari saat evaluasi akhir bulan/kuartal."
    ]
    for it in items_prob:
        p = tf_pr.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY
        p.space_before = Pt(12)

    # Kolom Kanan: Nilai Solusi ProMaP
    c_sol = add_card(slide2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), BLUE_LIGHT, RGBColor(191, 219, 254))
    tf_sol = c_sol.text_frame
    tf_sol.margin_left = tf_sol.margin_top = tf_sol.margin_right = tf_sol.margin_bottom = Inches(0.4)
    p_sol1 = tf_sol.paragraphs[0]
    p_sol1.text = "Nilai Inti yang Dihadirkan ProMaP V2.4"
    p_sol1.font.size = Pt(16)
    p_sol1.font.bold = True
    p_sol1.font.color.rgb = BLUE

    items_sol = [
        "Single Source of Truth: Seluruh divisi (Marketing, Ops, HR, Finance) bekerja dalam satu ekosistem terpadu.",
        "Akuntabilitas Berbasis Bukti (Evidence-Driven): Action Plan wajib melampirkan berkas bukti kerja sebelum disetujui Manager.",
        "Siklus 8 Status Transparan: Transisi status ketat dengan pencegahan manipulasi progres.",
        "Otomasi Overdue Tepat Waktu: Cron sistem otomatis menandai keterlambatan tepat pukul 00:01 WIB tanpa subjektivitas."
    ]
    for it in items_sol:
        p = tf_sol.add_paragraph()
        p.text = "✔ " + it
        p.font.size = Pt(11)
        p.font.color.rgb = NAVY
        p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: HIERARKI KERJA UNIVERSAL PROMAP
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Hierarki Kerja ProMaP: Dari Visi ke Aksi Lapangan", "STRUKTUR DATA & WORKFLOW")

    # Subtitle note
    sub_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.5))
    p_sub = sub_box.text_frame.paragraphs[0]
    p_sub.text = "Berlaku universal untuk semua divisi (Marketing, Operasional, HR, Keuangan, dsb), bukan hanya tim teknis."
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = TEXT_MUTED

    hier_steps = [
        ("1. PROJECT (INISIATIF)", "Payung Program Strategis", "Wadah besar inisiatif divisi, e.g. 'Ekspansi Pasar Q3', 'Audit SOP Gudang', 'Rekrutmen Batch 2'.", BLUE),
        ("2. TASK (PAKET SASARAN)", "Deliverable Terdelegasi", "Paket sasaran spesifik di bawah Project yang dipegang oleh PIC atau Lead penanggung jawab.", INDIGO),
        ("3. ACTION PLAN (EKSEKUSI)", "Aksi Harian / Mingguan", "Langkah kerja konkret PIC dengan target waktu, pembuktian (evidence), dan siklus 8 status.", EMERALD),
        ("4. CHECKLIST (MILESTONE)", "Verifikasi Detail", "Poin-poin mikro per Action Plan yang dapat dicentang langsung untuk mengukur kemajuan.", AMBER)
    ]

    hw = Inches(2.7)
    for idx, (title, subtitle, desc, col) in enumerate(hier_steps):
        h_left = Inches(0.8) + idx * (hw + Inches(0.3))
        card = add_card(slide3, h_left, Inches(2.3), hw, Inches(3.6))
        
        # Color bar top
        bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, h_left, Inches(2.3), hw, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(10)

    # 2 Feature Tambahan di Bawah
    c_bot1 = add_card(slide3, Inches(0.8), Inches(6.1), Inches(5.6), Inches(0.95), SLATE_100, BORDER_COLOR)
    tf_b1 = c_bot1.text_frame
    tf_b1.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "💡 Proposal (Bottom-Up Innovation)"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = BLUE
    p_b1_desc = tf_b1.add_paragraph()
    p_b1_desc.text = "Staf/PIC dapat mengusulkan inisiatif baru ke Manager. Jika disetujui, langsung jadi Action Plan."
    p_b1_desc.font.size = Pt(9.5)
    p_b1_desc.font.color.rgb = NAVY

    c_bot2 = add_card(slide3, Inches(6.8), Inches(6.1), Inches(5.7), Inches(0.95), SLATE_100, BORDER_COLOR)
    tf_b2 = c_bot2.text_frame
    tf_b2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b2 = tf_b2.paragraphs[0]
    p_b2.text = "📌 Personal Task (To-Do Pribadi)"
    p_b2.font.size = Pt(11)
    p_b2.font.bold = True
    p_b2.font.color.rgb = INDIGO
    p_b2_desc = tf_b2.add_paragraph()
    p_b2_desc.text = "Catatan tugas pribadi PIC untuk produktivitas mandiri tanpa mengotori agregasi progres tim."
    p_b2_desc.font.size = Pt(9.5)
    p_b2_desc.font.color.rgb = NAVY

    # =========================================================================
    # SLIDE 4: PETA ROLE & MATRIKS HAK AKSES (RBAC OVERVIEW)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Peta Peran (RBAC) & Hak Akses ProMaP", "ROLE & ACCESS CONTROL")

    role_matrix = [
        ("SUPER_ADMIN", "Super Admin", "Lintas Seluruh Perusahaan", "Kelola Perusahaan, Tier Langganan, Leads Demo, Global Log", BLUE),
        ("ADMIN_OPERATIONAL", "Admin Ops", "1 Perusahaan (Semua Divisi)", "Kelola Divisi, User Management, Kuota Kursi, Approval UserLabel", INDIGO),
        ("MANAGER", "Manager Divisi", "1 Divisi Penanggung Jawab", "Inisiatif Proyek, Review Action Plan, Approval Proposal, Workload Tim", EMERALD),
        ("PIC", "Person In Charge", "Pekerjaan Sendiri & Divisi", "Konsol 'My Work', Eksekusi Action Plan, Upload Evidence, Buat Proposal", AMBER),
        ("GUEST", "Tamu / Demo", "Data Simulasi (Dummy)", "Akses demo 2 jam, uji coba fitur interaktif, pendaftaran trial 30 hari", TEXT_MUTED)
    ]

    r_w = Inches(11.733)
    card_rh = Inches(0.88)
    for idx, (code, label, scope, rights, col) in enumerate(role_matrix):
        top_pos = Inches(1.8) + idx * (card_rh + Inches(0.18))
        card = add_card(slide4, Inches(0.8), top_pos, r_w, card_rh)
        
        # Color indicator tag
        tag = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.15), card_rh)
        tag.fill.solid()
        tag.fill.fore_color.rgb = col
        tag.line.fill.background()

        tf = card.text_frame
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.12)
        
        p1 = tf.paragraphs[0]
        p1.text = f"{label}  ({code})   •   Cakupan Data: {scope}"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"Wewenang Utama: {rights}"
        p2.font.size = Pt(10)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 5: DEEP DIVE: SUPER ADMIN & ADMIN OPERASIONAL
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Deep Dive: Super Admin & Admin Operasional", "TATA KELOLA TINGKAT TINGGI")

    # Super Admin Card
    c_sa = add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_sa = c_sa.text_frame
    tf_sa.margin_left = tf_sa.margin_top = tf_sa.margin_right = tf_sa.margin_bottom = Inches(0.35)
    p_sa = tf_sa.paragraphs[0]
    p_sa.text = "🛡️ Super Admin (Platform Owner)"
    p_sa.font.size = Pt(16)
    p_sa.font.bold = True
    p_sa.font.color.rgb = BLUE

    sa_points = [
        ("Multi-Tenant Governance", "Mengontrol pendaftaran perusahaan (Company CRUD), status langganan (Basic, Premium, Enterprise)."),
        ("Lead & Demo Management", "Memantau calon pelanggan dari formulir Guest Demo untuk dikonversi menjadi klien berbayar."),
        ("Global System Health", "Memantau performa Neon DB, integritas sistem, audit log aktivitas pengguna lintas tenant."),
        ("Keamanan Sentral", "Mengatur token rahasia cron, perlindungan isolasi tenant berbasis RLS logic.")
    ]
    for h, b in sa_points:
        p = tf_sa.add_paragraph()
        p.text = f"• {h}: {b}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = NAVY
        p.space_before = Pt(10)

    # Admin Ops Card
    c_ao = add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tf_ao = c_ao.text_frame
    tf_ao.margin_left = tf_ao.margin_top = tf_ao.margin_right = tf_ao.margin_bottom = Inches(0.35)
    p_ao = tf_ao.paragraphs[0]
    p_ao.text = "🏢 Admin Operasional (Company Admin)"
    p_ao.font.size = Pt(16)
    p_ao.font.bold = True
    p_ao.font.color.rgb = INDIGO

    ao_points = [
        ("Struktur Divisi", "Membuat dan mengelola seluruh divisi perusahaan (Marketing, Finance, HR, Operasional)."),
        ("Manajemen Pengguna & Kuota", "Menambah akun staf, menetapkan peran, serta membatasi seat quota sesuai paket langganan."),
        ("Persetujuan Label Jabatan", "Menyetujui usulan 'UserLabel' dinamis (e.g. Staff Tetap, Magang, Specialist, Lead)."),
        ("Pengawasan Lintas Divisi", "Menganalisis perbandingan completion rate dan tata kelola (Governance Index) seluruh divisi.")
    ]
    for h, b in ao_points:
        p = tf_ao.add_paragraph()
        p.text = f"• {h}: {b}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = NAVY
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 6: DEEP DIVE: MANAGER (ORKESTRATOR DIVISI)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Deep Dive: Manager (Penggerak & Quality Control)", "ORKESTRASI TIM DIVISI")

    mgr_cols = [
        ("Perencanaan Strategis", "Membuat Project (inisiatif besar) dan memecahnya menjadi Task deliverable serta mendelegasikan PIC yang tepat.", BLUE),
        ("Action Required Panel", "Mendapat notifikasi sentral di dashboard untuk Action Plan yang butuh review, approval, atau revisi segera.", AMBER),
        ("Approval Berbasis Bukti", "Memeriksa lampiran evidence PIC. Dapat menerima (Approve), meminta bukti baru, atau menolak (Reject) disertai alasan.", EMERALD),
        ("Monitoring Overdue & Beban", "Memantau keterlambatan tim secara proaktif, mengirim reminder cepat ke PIC, dan menjaga distribusi beban kerja.", RED)
    ]

    m_w = Inches(2.7)
    for idx, (title, desc, col) in enumerate(mgr_cols):
        m_left = Inches(0.8) + idx * (m_w + Inches(0.3))
        card = add_card(slide6, m_left, Inches(1.8), m_w, Inches(4.2))
        
        # Color badge
        bdg = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, m_left, Inches(1.8), m_w, Inches(0.12))
        bdg.fill.solid()
        bdg.fill.fore_color.rgb = col
        bdg.line.fill.background()

        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(12)

    # Bottom Banner for Manager
    c_mbot = add_card(slide6, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.85), EMERALD_LIGHT, RGBColor(167, 243, 208))
    tf_mb = c_mbot.text_frame
    tf_mb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_mb = tf_mb.paragraphs[0]
    p_mb.text = "🎯 Nilai Manager: Tidak lagi mengejar laporan manual via chat; seluruh progres & bukti tersaji otomatis dalam satu layar inspeksi."
    p_mb.font.size = Pt(11)
    p_mb.font.bold = True
    p_mb.font.color.rgb = RGBColor(6, 95, 70)

    # =========================================================================
    # SLIDE 7: DEEP DIVE: PIC (EKSEKUTOR) & GUEST DEMO
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Deep Dive: PIC (Pelaksana) & Mode Guest Demo", "EKSEKUSI LAPANGAN & PROSPEK")

    # PIC Card
    c_pic = add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_pic = c_pic.text_frame
    tf_pic.margin_left = tf_pic.margin_top = tf_pic.margin_right = tf_pic.margin_bottom = Inches(0.35)
    p_pic = tf_pic.paragraphs[0]
    p_pic.text = "⚡ PIC (Person In Charge / Eksekutor)"
    p_pic.font.size = Pt(16)
    p_pic.font.bold = True
    p_pic.font.color.rgb = AMBER

    pic_features = [
        ("Konsol Mandiri 'My Work'", "Fokus hanya pada tugas milik sendiri tanpa pusing administrasi organisasi besar."),
        ("Alur Bukti (Upload Evidence)", "Mengunggah file PDF, gambar, atau link hasil kerja nyata saat mengajukan approval."),
        ("Revisi Terarah", "Menerima catatan perbaikan spesifik dari Manager jika Action Plan ditolak (REJECTED)."),
        ("Inisiatif Usulan (Proposals)", "PIC dapat mengusulkan ide program kerja inovatif langsung ke Manager secara formal."),
        ("Multi-View Pribadi", "Tampilan Kanban Board, Tabel, dan Kalender yang sinkron dengan deadline pribadi.")
    ]
    for h, b in pic_features:
        p = tf_pic.add_paragraph()
        p.text = f"• {h}: {b}"
        p.font.size = Pt(10)
        p.font.color.rgb = NAVY
        p.space_before = Pt(8)

    # Guest Demo Card
    c_gst = add_card(slide7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    tf_gst = c_gst.text_frame
    tf_gst.margin_left = tf_gst.margin_top = tf_gst.margin_right = tf_gst.margin_bottom = Inches(0.35)
    p_gst = tf_gst.paragraphs[0]
    p_gst.text = "🌐 Mode Guest Demo (Eksplorasi Tamu)"
    p_gst.font.size = Pt(16)
    p_gst.font.bold = True
    p_gst.font.color.rgb = TEXT_MUTED

    gst_features = [
        ("Tanpa Registrasi Rumit", "Cukup mengisi Nama, Email, No HP, dan Nama Perusahaan langsung di card login."),
        ("Data Dummy Terisolasi", "Data demo terisi otomatis untuk eksplorasi tanpa menyentuh database operasional real."),
        ("Sesi Aktif 2 Jam", "Dibatasi token 2 jam dengan kuota maksimal 3x login demo per hari per email."),
        ("Funnel Leads B2B", "Setiap klik fitur aktif menghadirkan CTA 'Aktivasi Trial 30 Hari' yang langsung tercatat sebagai prospek tim sales.")
    ]
    for h, b in gst_features:
        p = tf_gst.add_paragraph()
        p.text = f"• {h}: {b}"
        p.font.size = Pt(10)
        p.font.color.rgb = NAVY
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 8: ALUR PENGGUNAAN 1: DARI USULAN KE AKSI (PROPOSALS)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Alur 1: Inovasi Bottom-Up Lewat Sistem Proposal", "ALUR PENGGUNAAN UTAMA")

    flow_prop = [
        ("1. SUBMIT PROPOSAL", "Staf / PIC membuat draft inisiatif, menjelaskan latar belakang, estimasi dampak, & anggaran.", INDIGO),
        ("2. ACTION REQUIRED", "Usulan masuk otomatis ke dashboard Manager divisi pada panel 'Action Required'.", AMBER),
        ("3. REVIEW & KEPUTUSAN", "Manager mengevaluasi usulan. Jika ditolak wajib menyertakan alasan penolakan tertulis.", BLUE),
        ("4. OTOMATIS JADI AP", "Jika disetujui, proposal dikonversi langsung menjadi Action Plan resmi di dalam Project.", EMERALD)
    ]

    pw = Inches(2.7)
    for idx, (title, desc, col) in enumerate(flow_prop):
        p_left = Inches(0.8) + idx * (pw + Inches(0.3))
        card = add_card(slide8, p_left, Inches(2.0), pw, Inches(4.5))
        
        # Step number pill
        pill = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_left + Inches(0.2), Inches(2.3), Inches(0.6), Inches(0.3))
        pill.fill.solid()
        pill.fill.fore_color.rgb = col
        pill.line.fill.background()
        p_num = pill.text_frame.paragraphs[0]
        p_num.text = str(idx + 1)
        p_num.font.bold = True
        p_num.font.size = Pt(11)
        p_num.font.color.rgb = WHITE
        p_num.alignment = PP_ALIGN.CENTER

        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.85)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(10)

    # =========================================================================
    # SLIDE 9: ALUR PENGGUNAAN 2: SIKLUS 8 STATUS ACTION PLAN
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Alur 2: Siklus 8 Status & Eksekusi Action Plan", "STANDAR OPERASIONAL EKSEKUSI")

    statuses = [
        ("NOT_STARTED", "Baru ditugaskan ke PIC, belum mulai dikerjakan.", SLATE_100, TEXT_MUTED),
        ("IN_PROGRESS", "Sedang dikerjakan oleh PIC secara aktif.", BLUE_LIGHT, BLUE),
        ("PENDING_APPROVAL", "Bukti kerja di-submit, menunggu review Manager.", AMBER_LIGHT, AMBER),
        ("EVIDENCE_REQUIRED", "Manager meminta bukti lampiran tambahan dari PIC.", AMBER_LIGHT, AMBER),
        ("APPROVED", "Disetujui Manager setelah bukti terverifikasi valid.", EMERALD_LIGHT, EMERALD),
        ("COMPLETE", "Pekerjaan tuntas sempurna dan tercatat di histori capaian.", EMERALD_LIGHT, EMERALD),
        ("REJECTED", "Ditolak karena tidak sesuai SOP; kembali ke In Progress untuk revisi.", RGBColor(254, 242, 242), RED),
        ("OVERDUE", "Terlambat otomatis oleh Cron 00:01 WIB jika lewat tenggat.", RGBColor(254, 242, 242), RED)
    ]

    card_s_w = Inches(5.6)
    card_s_h = Inches(1.0)
    for idx, (code, desc, bg_c, text_c) in enumerate(statuses):
        col_idx = idx % 2
        row_idx = idx // 2
        s_left = Inches(0.8) + col_idx * (card_s_w + Inches(0.533))
        s_top = Inches(1.8) + row_idx * (card_s_h + Inches(0.2))

        c = add_card(slide9, s_left, s_top, card_s_w, card_s_h, bg_c, BORDER_COLOR)
        tf = c.text_frame
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        
        p1 = tf.paragraphs[0]
        p1.text = code
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = text_c

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 10: FLEKSIBILITAS MULTI-VIEW WORKSPACE
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Multi-View Workspace: 1 Data, 3 Cara Pandang", "USER EXPERIENCE & TAMPILAN")

    views = [
        ("📋 Table View (Tabel Detail)", "Inspeksi & Filter Mendalam", 
         "Cocok untuk audit menyeluruh:\n• Filter canggih: Divisi, PIC, Prioritas, Periode\n• Sortir cepat deadline terdekat & prioritas\n• Drawer samping untuk inspeksi tanpa pindah halaman", BLUE),
        ("📌 Kanban Board (5 Kolom)", "Visibilitas Progres Cepat", 
         "Memetakan 8 status ke dalam 5 kolom rapi:\n• Not Started, In Progress, Review, Needs Revision, Done\n• Badge merah Overdue di kartu\n• Drag & drop terlindungi aturan status flow", INDIGO),
        ("📅 Calendar & Gantt Timeline", "Manajemen Batas Waktu", 
         "Visualisasi horizontal rentang waktu:\n• Tampilan bulanan, mingguan, & agenda\n• Ekspor kalender ke format PDF resolusi tinggi\n• Sinkronisasi feed kalender (.ics)", EMERALD)
    ]

    vw = Inches(3.7)
    for idx, (title, sub, details, col) in enumerate(views):
        v_left = Inches(0.8) + idx * (vw + Inches(0.3))
        card = add_card(slide10, v_left, Inches(1.8), vw, Inches(5.0))
        
        # Color bar
        bar = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, v_left, Inches(1.8), vw, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(4)

        for line in details.split("\n"):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(10)
            p.font.color.rgb = NAVY
            p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 11: MONITORING, REPORTING & GOVERNANCE
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Sistem Monitoring, Notifikasi & Tata Kelola", "PEMANTAUAN & LAPORAN")

    gov_cards = [
        ("🔔 Notifikasi & @Mention", "Pemberitahuan in-app instan saat Action Plan didelegasikan, butuh review, atau saat staf di-mention dalam diskusi."),
        ("⏰ Cron Job Overdue Otomatis", "Berjalan setiap pukul 00:01 WIB. Menandai tugas yang melewati deadline tanpa intervensi manual."),
        ("📊 Dashboard Metrik Adaptif", "Menampilkan 4 kartu ringkasan (Total, Selesai, In Progress, Overdue) serta Donut chart sebaran status."),
        ("📈 Laporan Eksekutif & Audit Log", "Ekspor laporan evaluasi berkala (Excel/PDF) dan riwayat perubahan lengkap (siapa mengubah apa dan kapan).")
    ]

    cw = Inches(5.6)
    ch = Inches(2.2)
    for idx, (title, desc) in enumerate(gov_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        g_left = Inches(0.8) + col_idx * (cw + Inches(0.533))
        g_top = Inches(2.0) + row_idx * (ch + Inches(0.4))

        card = add_card(slide11, g_left, g_top, cw, ch)
        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.35)
        tf.margin_top = Inches(0.25)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = BLUE

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 12: KESIMPULAN & NILAI TAMBAH BISNIS
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Nilai Tambah Nyata ProMaP bagi Perusahaan", "RINGKASAN & DAMPAK")

    impacts = [
        ("Untuk Pimpinan / Direksi", "Transparansi penuh terhadap kemajuan inisiatif strategis seluruh divisi tanpa menunggu laporan akhir bulan.", BLUE),
        ("Untuk Manager Divisi", "Kontrol kualitas kerja yang objektif dengan bukti riil, delegasi tugas presisi, dan deteksi dini hambatan.", INDIGO),
        ("Untuk Pelaksana (PIC)", "Kejelasan prioritas kerja harian di konsol 'My Work', tidak ada lagi miskomunikasi target atau tugas ganda.", EMERALD)
    ]

    iw = Inches(3.7)
    for idx, (title, desc, col) in enumerate(impacts):
        i_left = Inches(0.8) + idx * (iw + Inches(0.3))
        card = add_card(slide12, i_left, Inches(1.8), iw, Inches(3.8))
        
        bar = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, i_left, Inches(1.8), iw, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tf = card.text_frame
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.3)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = NAVY
        p2.space_before = Pt(12)

    # Final CTA Card
    c_final = add_card(slide12, Inches(0.8), Inches(5.9), Inches(11.733), Inches(1.1), NAVY, SLATE_800)
    tf_f = c_final.text_frame
    tf_f.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_f1 = tf_f.paragraphs[0]
    p_f1.text = "Siap Mengubah Rencana Strategis Menjadi Tindakan Nyata Berakuntabilitas Tinggi"
    p_f1.font.size = Pt(13)
    p_f1.font.bold = True
    p_f1.font.color.rgb = WHITE
    p_f1.alignment = PP_ALIGN.CENTER

    p_f2 = tf_f.add_paragraph()
    p_f2.text = "ProMaP Enterprise V2.4 — Satu Workspace untuk Seluruh Divisi Perusahaan"
    p_f2.font.size = Pt(10.5)
    p_f2.font.color.rgb = RGBColor(148, 163, 184)
    p_f2.alignment = PP_ALIGN.CENTER
    p_f2.space_before = Pt(4)

    output_path = os.path.abspath("Presentasi_Alur_dan_Role_ProMaP_V2.pptx")
    prs.save(output_path)
    print(f"SUCCESS: {output_path}")

if __name__ == "__main__":
    create_deck()
